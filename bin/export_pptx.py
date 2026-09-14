#!/usr/bin/env python3
"""Export a rendered talk as a PowerPoint file: one slide per step, the step's clip filling the slide and starting
automatically when the slide appears, the step's poster frame shown before it plays, and the speaker note in the
slide's notes. The video holds on its last frame, which is the frame the presenter would hold on.

Why this shape: PowerPoint cannot pause one long video at arbitrary points from a click, so the per-step clips that
Manim writes with --save_sections are the natural unit; the JSON index gives their order and durations. The autoplay
is the timing tree PowerPoint itself writes for a video set to "Start: Automatically"; python-pptx has no API for it,
so the XML is inserted after the movie is added.

Usage: bin/export_pptx.py <talk> [ql|qm|qh] [out.pptx] [--click]   default quality qh, output <talk>/<talk>.pptx;
--click leaves the clips to start on click instead of automatically
Requires python-pptx (pip install python-pptx) and ffmpeg (poster frames)."""
import glob, json, os, re, subprocess, sys, tempfile
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from lib import theme as _theme
from lib.talks import dir_of, quality, rendered_or_exit, scenes_of

CLICK = "--click" in sys.argv
sys.argv = [a for a in sys.argv if a != "--click"]
talk = sys.argv[1] if len(sys.argv) > 1 else sys.exit(__doc__)
Q = quality(sys.argv[2] if len(sys.argv) > 2 else None)
root = dir_of(talk)
out = sys.argv[3] if len(sys.argv) > 3 else f"{root}/{talk}.pptx"
items = rendered_or_exit(root, Q, talk)
# The slide behind the clip is the talk's own ground, resolved as the render resolves it: a bright deck exported on its
# own must not be letterboxed in near-black.
BG_HEX = _theme.load(_theme.active_name(root))["bg"].lstrip("#")

title_line = next((l for l in open(f"{root}/script.md") if l.startswith("# ")), f"# {talk}") if os.path.exists(f"{root}/script.md") else f"# {talk}"
TITLE = title_line[2:].strip()

NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
AUTOPLAY = """<p:timing xmlns:p="%s"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="withEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="%d" fill="hold"/><p:tgtEl><p:spTgt spid="%d"/></p:tgtEl></p:cBhvr></p:cmd>
</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn>
<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>
</p:seq>
<p:video><p:cMediaNode vol="80000"><p:cTn id="7" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="%d"/></p:tgtEl></p:cMediaNode></p:video>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>"""


def poster(clip: str, png: str):
    subprocess.run(["ffmpeg", "-loglevel", "error", "-y", "-i", clip, "-frames:v", "1", "-vf", "scale=1280:-1", png], check=True)


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]
for stem, scene in scenes_of(root):
    if not any(scene == sc for _, sc, _, _ in items):
        print(f"not rendered: {scene}")
n_steps = 0
with tempfile.TemporaryDirectory() as tmp:            # the poster frames are throwaway; leaving them behind cost megabytes a run
    for stem, scene, _video, index in items:
        notes_path = f"{root}/media/notes/{scene}.json"
        notes = json.load(open(notes_path)) if os.path.exists(notes_path) else []
        for k, sec in enumerate(json.load(open(index))):
            clip = f"{root}/media/videos/{stem}/{Q}/sections/{sec['video']}"
            png = f"{tmp}/{scene}-{k}.png"
            poster(clip, png)
            slide = prs.slides.add_slide(blank)
            try:                                     # the clip is letterboxed on a 16:9 slide; the margins must not flash white
                slide.background.fill.solid()
                slide.background.fill.fore_color.rgb = RGBColor.from_string(BG_HEX)
            except Exception:                        # an older python-pptx cannot set a slide background; the clip still fills it
                pass
            movie = slide.shapes.add_movie(clip, 0, 0, prs.slide_width, prs.slide_height, poster_frame_image=png, mime_type="video/mp4")
            spid = movie.shape_id
            if not CLICK:
                timing = etree.fromstring(AUTOPLAY % (NS, int(float(sec["duration"]) * 1000), spid, spid))
                old = slide.element.find(f"{{{NS}}}timing")   # python-pptx already wrote a bare media timing node for the movie
                if old is not None:                           # two p:timing elements make the file invalid; replace, do not append
                    slide.element.replace(old, timing)
                else:
                    slide.element.append(timing)
            note = notes[k] if k < len(notes) else ""
            slide.notes_slide.notes_text_frame.text = f"{scene}, step {k + 1}\n\n{note}"
            n_steps += 1
prs.core_properties.title = TITLE
prs.save(out)
print(f"{talk}: {n_steps} slides -> {out}")
